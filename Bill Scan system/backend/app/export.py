# backend/app/export.py
"""
Spreadsheet exports.

Two shapes of the same data:

  bills  one row per bill, for totalling and pivoting
  lines  one row per line item, carrying its bill's date and vendor so the
         sheet stands on its own without a lookup

Both the .xlsx download and the CSV endpoint Excel binds to are built from
these, so a column added here appears in both.
"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Bill

# (header, attribute-or-callable, number format). The format is applied in the
# .xlsx; CSV writes the raw value.
MONEY = "#,##0.00"
DATE = "yyyy-mm-dd"

BILL_COLUMNS = [
    ("Bill ID", lambda b: b.id, None),
    ("Date", lambda b: b.bill_date, DATE),
    ("Vendor", lambda b: b.vendor_name, None),
    ("GSTIN", lambda b: b.vendor_gstin, None),
    ("Bill No", lambda b: b.bill_number, None),
    ("Category", lambda b: (b.category or "").replace("_", " "), None),
    ("Paid By", lambda b: b.payment_method, None),
    ("Subtotal", lambda b: b.subtotal, MONEY),
    ("CGST", lambda b: b.cgst, MONEY),
    ("SGST", lambda b: b.sgst, MONEY),
    ("IGST", lambda b: b.igst, MONEY),
    ("Tax", lambda b: b.tax_amount, MONEY),
    ("Round Off", lambda b: b.round_off, MONEY),
    ("Grand Total", lambda b: b.grand_total, MONEY),
    ("Items", lambda b: len(b.line_items or []), None),
    ("Status", lambda b: b.status, None),
    ("Entered By", lambda b: "hand" if b.source == "manual" else "AI", None),
    ("Confidence", lambda b: b.effective_confidence, None),
    ("Math OK", lambda b: "yes" if b.validation_ok else "NO", None),
    ("Flags", lambda b: _flag_summary(b), None),
]

LINE_COLUMNS = [
    ("Bill ID", lambda b, ln, i: b.id, None),
    ("Date", lambda b, ln, i: b.bill_date, DATE),
    ("Vendor", lambda b, ln, i: b.vendor_name, None),
    ("Category", lambda b, ln, i: (b.category or "").replace("_", " "), None),
    ("Line", lambda b, ln, i: i + 1, None),
    ("Item", lambda b, ln, i: ln.get("description"), None),
    ("Qty", lambda b, ln, i: ln.get("quantity"), "#,##0.###"),
    ("Unit", lambda b, ln, i: ln.get("unit"), None),
    ("Rate", lambda b, ln, i: ln.get("rate"), MONEY),
    ("Amount", lambda b, ln, i: ln.get("amount"), MONEY),
    ("Tax %", lambda b, ln, i: ln.get("tax_percent"), None),
]


def _flag_summary(bill: Bill) -> str:
    """The field paths that failed, so a filtered sheet shows what to chase."""
    flags = bill.validation_flags or []
    errors = [f["field_path"] for f in flags if f.get("severity") == "error"]
    return ", ".join(errors)


def bill_rows(bills: list[Bill]) -> tuple[list[str], list[list]]:
    headers = [c[0] for c in BILL_COLUMNS]
    rows = [[fn(b) for _, fn, _ in BILL_COLUMNS] for b in bills]
    return headers, rows


def line_rows(bills: list[Bill]) -> tuple[list[str], list[list]]:
    headers = [c[0] for c in LINE_COLUMNS]
    rows = [
        [fn(b, line, i) for _, fn, _ in LINE_COLUMNS]
        for b in bills
        for i, line in enumerate(b.line_items or [])
    ]
    return headers, rows


def to_csv(bills: list[Bill], granularity: str = "bills") -> bytes:
    """
    CSV for the live Excel link.

    Encoded UTF-8 with a BOM: without it Excel reads the file as the system
    codepage and mangles any non-ASCII vendor name.
    """
    headers, rows = (
        line_rows(bills) if granularity == "lines" else bill_rows(bills)
    )
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(bills: list[Bill]) -> bytes:
    """A workbook with a Bills sheet and a Line Items sheet."""
    wb = Workbook()

    bills_ws = wb.active
    bills_ws.title = "Bills"
    _fill_sheet(bills_ws, *bill_rows(bills), BILL_COLUMNS)
    _add_total_row(bills_ws, len(bills), BILL_COLUMNS)

    lines_ws = wb.create_sheet("Line Items")
    _fill_sheet(lines_ws, *line_rows(bills), LINE_COLUMNS)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fill_sheet(ws, headers, rows, columns) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="374151")

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(row)

    for idx, (header, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        # Width from the widest value, so nothing shows as ####.
        widest = max(
            [len(str(header))]
            + [len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None]
            or [0]
        )
        ws.column_dimensions[letter].width = min(max(widest + 2, 9), 42)

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )


def _add_total_row(ws, count: int, columns) -> None:
    """A SUM under each money column, one blank row below the data."""
    if not count:
        return
    row = count + 3
    ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
    for idx, (_, _, fmt) in enumerate(columns, start=1):
        if fmt != MONEY:
            continue
        letter = get_column_letter(idx)
        cell = ws.cell(row=row, column=idx)
        cell.value = f"=SUM({letter}2:{letter}{count + 1})"
        cell.number_format = MONEY
        cell.font = Font(bold=True)
