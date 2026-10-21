# backend/tools/api_smoke_test.py
"""
End-to-end check of the HTTP surface with the vision call stubbed out.

Exercises scan -> review -> revalidate -> confirm -> train -> list against a
temporary database, so it runs offline and costs nothing. Run from backend/:

    python -m tools.api_smoke_test
"""

import dataclasses
import io
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Point the app at a throwaway database before anything imports settings.
_tmp = tempfile.mkdtemp(prefix="daily-bills-smoke-")
os.environ.setdefault("BILLS_MODEL", "claude-opus-5")

from PIL import Image  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import models  # noqa: E402
from app import config  # noqa: E402

# Rebind storage to the temp dir so the real data/ directory is never touched.
# Settings is frozen, so this replaces the instance rather than mutating it.
_uploads = Path(_tmp) / "uploads"
_uploads.mkdir(parents=True, exist_ok=True)
config.settings = dataclasses.replace(
    config.settings, upload_dir=_uploads, db_path=Path(_tmp) / "smoke.db"
)

models.engine.dispose()
models.engine = models.create_engine(
    f"sqlite:///{Path(_tmp) / 'smoke.db'}", connect_args={"check_same_thread": False}
)
models.SessionLocal.configure(bind=models.engine)
models.init_db()  # TestClient only runs the lifespan hook as a context manager

from app.main import app  # noqa: E402
from app.routers import bills as bills_router  # noqa: E402

failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  pass  {label}")
    else:
        print(f"  FAIL  {label} {detail}")
        failures.append(label)


# --- the extraction the stub returns --------------------------------------
EXTRACTED = {
    "error": None,
    "error_reason": None,
    "vendor_name": "ABC Store",
    "vendor_gstin": None,
    "bill_date": "2026-08-24",
    "bill_date_raw": None,
    "bill_number": "R-118",
    "payment_method": "cash",
    "category": "groceries",
    "confidence": "high",
    "line_items": [
        {"description": "Milk", "quantity": 2, "unit": "ltr",
         "rate": 30, "amount": 60, "tax_percent": None},
        {"description": "Bread", "quantity": 1, "unit": "pcs",
         "rate": 40, "amount": 40, "tax_percent": None},
    ],
    "subtotal": 100,
    "cgst": None,
    "sgst": None,
    "igst": None,
    "tax_amount": 0,
    "round_off": 0,
    "grand_total": 1000,  # a misread digit: the validator must catch it
    "reconciliation_warning": True,
}

captured_profiles = {}


def fake_extract(raw_image, vendor_profiles=None):
    captured_profiles["last"] = vendor_profiles
    return dict(EXTRACTED), {"input_tokens": 1200, "output_tokens": 300,
                             "cache_read_input_tokens": 0, "model": "stub"}


bills_router.extract_bill = fake_extract

client = TestClient(app)


def jpeg_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (600, 900), (250, 250, 250)).save(buf, format="JPEG")
    return buf.getvalue()


print("health")
r = client.get("/api/health")
check("200", r.status_code == 200, r.text)

print("upload rejects a non-image")
r = client.post("/api/bills", files={"file": ("notes.txt", b"hello", "text/plain")})
check("415", r.status_code == 415, r.text)

print("upload a bill")
r = client.post("/api/bills", files={"file": ("bill.jpg", jpeg_bytes(), "image/jpeg")})
check("201", r.status_code == 201, r.text)
bill = r.json()
bill_id = bill["id"]
check("vendor read", bill["vendor_name"] == "ABC Store")
check("misread total flagged", bill["_validation_ok"] is False)
check("grand_total is the flagged field",
      any(f["field_path"] == "grand_total" for f in bill["_validation_flags"]))
check("confidence downgraded to low", bill["_effective_confidence"] == "low")
check("model said high", bill["confidence"] == "high")
check("no vendors known yet", captured_profiles["last"] == [])

print("the original image comes back")
r = client.get(f"/api/bills/{bill_id}/image")
check("200", r.status_code == 200)
check("is a jpeg", r.headers["content-type"] == "image/jpeg")

print("revalidating the corrected bill clears the flags")
fixed = {k: v for k, v in bill.items()
         if not k.startswith("_") and k not in ("id", "status", "created_at")}
fixed["grand_total"] = 100
r = client.post(f"/api/bills/{bill_id}/validate", json=fixed)
check("200", r.status_code == 200, r.text)
check("now ok", r.json()["ok"] is True, r.text)
check("no flags", r.json()["flags"] == [])

print("revalidating an unknown bill is a 404")
r = client.post("/api/bills/9999/validate", json=fixed)
check("404", r.status_code == 404)

print("confirm and train the vendor")
r = client.post(
    f"/api/bills/{bill_id}/confirm",
    json={"bill": fixed, "train": True, "vendor_notes": "Corner shop, no GST"},
)
check("200", r.status_code == 200, r.text)
body = r.json()
check("status is confirmed", body["bill"]["status"] == "confirmed")
check("edit recorded", "grand_total" in body["edited_fields"], body["edited_fields"])
check("profile created", body["vendor_profile"]["vendor_name"] == "ABC Store")
check("items remembered",
      set(body["vendor_profile"]["common_items"]) == {"Milk", "Bread"})
check("note stored", body["vendor_profile"]["notes"] == "Corner shop, no GST")

print("the raw extraction is kept for later quality measurement")
with models.SessionLocal() as db:
    row = db.get(models.Bill, bill_id)
    check("raw grand_total preserved", row.raw_extraction["grand_total"] == 1000)
    check("stored grand_total corrected", row.grand_total == 100)

print("the next scan carries the trained roster into the prompt")
r = client.post("/api/bills", files={"file": ("bill2.jpg", jpeg_bytes(), "image/jpeg")})
check("201", r.status_code == 201, r.text)
roster = captured_profiles["last"]
check("roster has the vendor", [p["vendor_name"] for p in roster] == ["ABC Store"])
check("roster carries the items", "Milk" in roster[0]["common_items"])

print("a vendor name variant maps to the same profile")
check("key normalised",
      bills_router.vendor_key("A.B.C.  STORE ") == bills_router.vendor_key("ABC Store"))

print("listing and filtering")
r = client.get("/api/bills")
check("200", r.status_code == 200)
check("two bills", len(r.json()["bills"]) == 2, r.text)
r = client.get("/api/bills?status=confirmed")
check("one confirmed", len(r.json()["bills"]) == 1)
r = client.get("/api/bills?category=fuel")
check("none in fuel", len(r.json()["bills"]) == 0)

print("vendors endpoint")
r = client.get("/api/vendors")
check("one vendor", len(r.json()["vendors"]) == 1, r.text)

print("a bill can be entered by hand, with no API call at all")
r = client.post("/api/bills/manual")
check("201", r.status_code == 201, r.text)
manual = r.json()
mid = manual["id"]
check("marked manual", manual["_source"] == "manual")
check("no photo yet", manual["_has_image"] is False)
check("starts with one blank line", len(manual["line_items"]) == 1)
check("blank bill does not reconcile", manual["_validation_ok"] is False)

print("a photo-less bill reports no photo, and deletes cleanly")
# Path("") is ".", so an unguarded unlink here tries to delete the working
# directory rather than a file.
r = client.post("/api/bills/manual")
throwaway = r.json()["id"]
check("image is a 404, not a crash",
      client.get(f"/api/bills/{throwaway}/image").status_code == 404)
check("deletes without touching the cwd",
      client.delete(f"/api/bills/{throwaway}").status_code == 204)

print("typing it in reconciles it")
typed = {
    "vendor_name": "Hand Entry Shop",
    "bill_date": "2026-08-24",
    "payment_method": "cash",
    "category": "vegetables",
    "line_items": [
        {"description": "Onions", "quantity": 3, "unit": "kg",
         "rate": 25, "amount": 75, "tax_percent": None},
    ],
    "subtotal": 75, "tax_amount": 0, "round_off": 0, "grand_total": 75,
    "confidence": "medium",
}
r = client.post(f"/api/bills/{mid}/validate", json=typed)
check("validates clean", r.json()["ok"] is True, r.text)

print("a photo can be attached afterwards")
r = client.post(
    f"/api/bills/{mid}/image",
    files={"file": ("paper.jpg", jpeg_bytes(), "image/jpeg")},
)
check("200", r.status_code == 200, r.text)
check("now has an image", r.json()["_has_image"] is True)
check("image is served", client.get(f"/api/bills/{mid}/image").status_code == 200)

print("a hand-entered bill confirms and trains like any other")
r = client.post(f"/api/bills/{mid}/confirm", json={"bill": typed, "train": True})
check("200", r.status_code == 200, r.text)
check("confirmed", r.json()["bill"]["status"] == "confirmed")
check("vendor trained", r.json()["vendor_profile"]["vendor_name"] == "Hand Entry Shop")
check("still marked manual", r.json()["bill"]["_source"] == "manual")

print("an attached photo replaces the old one without orphaning it")
before = len(list(_uploads.glob("*.jpg")))
r = client.post(
    f"/api/bills/{mid}/image",
    files={"file": ("paper2.jpg", jpeg_bytes(), "image/jpeg")},
)
check("200", r.status_code == 200)
check("old file removed", len(list(_uploads.glob("*.jpg"))) == before,
      f"{before} -> {len(list(_uploads.glob('*.jpg')))}")

print("a server with no credentials says so instead of returning a 500")
# The real extractor, with a client that has no key: the SDK raises a bare
# TypeError while building the request, which must not escape as a 500.
import anthropic  # noqa: E402
from app.extraction import claude_vision  # noqa: E402

claude_vision._client = anthropic.Anthropic(api_key=None, auth_token=None)
bills_router.extract_bill = claude_vision.extract_bill
try:
    r = client.post("/api/bills", files={"file": ("b.jpg", jpeg_bytes(), "image/jpeg")})
    check("503, not 500", r.status_code == 503, f"got {r.status_code}")
    check("message names the key",
          "ANTHROPIC_API_KEY" in r.json().get("detail", ""), r.text)
except Exception as e:  # noqa: BLE001
    check("no unhandled exception", False, repr(e))
finally:
    bills_router.extract_bill = fake_extract
    claude_vision._client = None

print("CSV export, the endpoint Excel links to")
r = client.get("/api/bills.csv")
check("200", r.status_code == 200, r.text)
check("is csv", r.headers["content-type"].startswith("text/csv"))
check("has a BOM so Excel reads UTF-8", r.content.startswith(b"\xef\xbb\xbf"))
text = r.content.decode("utf-8-sig")
check("header row", text.splitlines()[0].startswith("Bill ID,Date,Vendor"))
check("CRLF line endings", "\r\n" in text)
expected = len(client.get("/api/bills").json()["bills"])
check("one row per bill, matching the list view",
      len(text.strip().splitlines()) == expected + 1,
      f"{len(text.strip().splitlines())} lines for {expected} bills")
check("served inline, not as a download",
      "attachment" not in r.headers.get("content-disposition", ""))

print("CSV honours the same filters as the list view")
r = client.get("/api/bills.csv?status=confirmed")
expected = len(client.get("/api/bills?status=confirmed").json()["bills"])
check("confirmed rows match the filtered list view",
      len(r.content.decode("utf-8-sig").strip().splitlines()) == expected + 1)
r = client.get("/api/bills.csv?category=fuel")
check("header only when nothing matches",
      len(r.content.decode("utf-8-sig").strip().splitlines()) == 1)

print("CSV at line granularity")
r = client.get("/api/bills.csv?granularity=lines")
lines = r.content.decode("utf-8-sig").strip().splitlines()
check("header names the item column", "Item" in lines[0])
check("a row per line item", len(lines) > 3, str(len(lines)))
r = client.get("/api/bills.csv?granularity=nonsense")
check("bad granularity rejected", r.status_code == 422)

print("Excel export")
r = client.get("/api/bills.xlsx")
check("200", r.status_code == 200)
check("is a workbook", r.content.startswith(b"PK"))
check("downloads as a file",
      "attachment" in r.headers.get("content-disposition", ""))
check("filename is dated", ".xlsx" in r.headers["content-disposition"])

from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(io.BytesIO(r.content))
check("two sheets", wb.sheetnames == ["Bills", "Line Items"], str(wb.sheetnames))
ws = wb["Bills"]
check("header is frozen", ws.freeze_panes == "A2")
check("has an autofilter", ws.auto_filter.ref is not None)
headers = [c.value for c in ws[1]]
check("Grand Total column present", "Grand Total" in headers)
check("flags column present", "Flags" in headers)
totals_col = headers.index("Grand Total") + 1
check("money cells are formatted",
      ws.cell(row=2, column=totals_col).number_format == "#,##0.00")
total_row = ws.max_row
check("a SUM total row",
      str(ws.cell(row=total_row, column=totals_col).value).startswith("=SUM("),
      repr(ws.cell(row=total_row, column=totals_col).value))
check("line items sheet has rows", wb["Line Items"].max_row > 1)

print("an empty export does not crash")
r = client.get("/api/bills.xlsx?category=fuel")
check("200", r.status_code == 200)
check("still a workbook", load_workbook(io.BytesIO(r.content))["Bills"].max_row == 1)

print("delete removes the row and the image")
r = client.delete(f"/api/bills/{bill_id}")
check("204", r.status_code == 204)
check("gone", client.get(f"/api/bills/{bill_id}").status_code == 404)

print()
if failures:
    print(f"{len(failures)} check(s) failed:")
    for f in failures:
        print(f"  - {f}")
    sys.exit(1)
print("api works")
